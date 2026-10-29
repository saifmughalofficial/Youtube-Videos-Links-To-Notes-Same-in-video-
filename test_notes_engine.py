"""Offline checks — no network, no API quota used.  Run:  python -m pytest -q"""
import io
import json
import os
import tempfile
from pathlib import Path

import pytest

os.environ["NOTES_CACHE_DIR"] = tempfile.mkdtemp(prefix="l2n_test_")
for _k in ("NOTES_API_KEY", "GEMINI_API_KEY", "GOOGLE_API_KEY", "NOTES_MODEL", "GROQ_API_KEY", "GROQ_MODEL",
           "OPENROUTER_API_KEY", "OPENROUTER_MODEL"):
    os.environ.pop(_k, None)

import notes_engine as nt  # noqa: E402

HERE = Path(__file__).parent

SAMPLE = {
    "title": "Bus fare revenue — demand function", "video_id": "abcdefghijk", "duration": 1500,
    "max_height": 360, "created": "28 Oct 2026, 22:00", "lectures": 1,
    "items": [
        {"type": "heading", "text": "Business Mathematics"},
        {"type": "question", "text": "Problem, page 778: The transit authority has surveyed citizens…\n"
                                     "(a) Determine the fare which should be charged to maximize revenue.\n"
                                     "(b) What is the expected maximum revenue?"},
        {"type": "step", "text": "q = 10,000 − 125p → ①"},
        {"type": "step", "text": "R = pq → ②"},
        {"type": "step", "text": "R = 10,000p − 125p²"},
        {"type": "answer", "text": "p = 40 cents"},
        {"type": "formula", "text": "dR/dp = 0"},
        {"type": "question", "text": "Find the maximum value of R(p) = −10p² + 1750p"},
        {"type": "step", "text": "R′(p) = −20p + 1750 = 0 ⇒ p = 87.5"},
        {"type": "answer", "text": "R = 76,562.5"},
    ],
}


@pytest.fixture(autouse=True)
def fast(monkeypatch):
    import time as real_time
    monkeypatch.setattr(nt, "MIN_GAP_SECONDS", 0)
    monkeypatch.setattr(nt, "time", type("T", (), {"time": staticmethod(real_time.time),
                                                    "sleep": staticmethod(lambda *_: None)}))
    nt._dead.clear()
    nt._last_call.clear()


# ------------------------------------------------------------------ links

@pytest.mark.parametrize("link", [
    "https://www.youtube.com/watch?v=qaMHTSugMUA", "https://youtu.be/qaMHTSugMUA?si=xyz",
    "youtube.com/shorts/qaMHTSugMUA", "https://youtube.com/shorts/qaMHTSugMUA?feature=share",
    "https://m.youtube.com/watch?feature=share&v=qaMHTSugMUA", "https://www.youtube.com/live/qaMHTSugMUA?t=10",
    "https://www.youtube.com/embed/qaMHTSugMUA", "qaMHTSugMUA",
])
def test_parse_video_id(link):
    assert nt.parse_video_id(link) == "qaMHTSugMUA"


def test_bad_links():
    assert nt.parse_video_id("https://vimeo.com/123") is None
    assert nt.parse_video_id("hello") is None
    assert nt.parse_playlist_id("https://youtube.com/playlist?list=PLabcdefghij12") == "PLabcdefghij12"


def test_split_links_from_a_numbered_list():
    text = ("Basic Shorts::: Derivatives\n1\nhttps://youtube.com/shorts/YhR0ESQpTTY?feature=share\n2\n"
            "https://youtube.com/shorts/-RjJuuAb3AI?feature=share\n23\nhttps://youtube.com/shorts/JcqIZo68ags?feature=share\n"
            "24\nhttps://youtube.com/shorts/JcqIZo68ags?feature=share\n\nAssignment 1 https://youtu.be/qaMHTSugMUA")
    links = nt.split_links(text)
    assert [nt.parse_video_id(x) for x in links] == ["YhR0ESQpTTY", "-RjJuuAb3AI", "JcqIZo68ags", "qaMHTSugMUA"]


# ------------------------------------------------------------------ cleaning

@pytest.mark.parametrize("raw,want", [
    (r"$x^2 + y^{10}$", "x² + y¹⁰"),
    (r"\frac{a+b}{c}", "(a+b)/(c)"),
    (r"\sqrt{x+1} \leq 3", "√(x+1) ≤ 3"),
    (r"a_1 + a_{12}", "a₁ + a₁₂"),
    (r"2 \times 3 \div 4 \pm 1", "2 × 3 ÷ 4 ± 1"),
    (r"\frac{3}{4}", "3/4"),
    (r"x^{n+1}", "xⁿ⁺¹"),
    (r"e^{2x}", "e^(2x)"),
])
def test_tidy_math(raw, want):
    assert nt.tidy_math(raw) == want


def test_clean_items():
    items = nt.clean_items([
        {"type": "question", "text": "Q.3: Find x if 2x = 6 www.maths.com"},
        {"type": "question", "text": "Quadratic formula derivation"},
        {"type": "answer", "text": "Ans: x = 3"},
        {"type": "step", "text": "Step 2) x = 6/2"},
        {"type": "text", "text": "Please subscribe to my channel"},
        {"type": "solution", "text": "2x = 6 ہے"},
        {"type": "note", "text": "spoken tip becomes plain text"},
        "not a dict",
    ])
    assert [i["text"] for i in items] == ["Find x if 2x = 6", "Quadratic formula derivation", "x = 3",
                                          "x = 6/2", "2x = 6", "spoken tip becomes plain text"]
    assert [i["type"] for i in items] == ["question", "question", "answer", "step", "step", "text"]


def test_join_parts_drops_overlap_but_keeps_real_repeats():
    a = [{"type": "question", "text": "Solve 2x = 4"}, {"type": "step", "text": "x = 2"},
         {"type": "answer", "text": "x = 2"}]
    b = [{"type": "answer", "text": "x = 2"}, {"type": "question", "text": "Solve 3x = 6"},
         {"type": "step", "text": "x = 2"}, {"type": "answer", "text": "x = 2"}]
    assert [i["text"] for i in nt.join_parts(a, b)] == ["Solve 2x = 4", "x = 2", "x = 2", "Solve 3x = 6", "x = 2", "x = 2"]


def test_extract_json_repairs_cut_off_answer():
    assert nt.extract_json('```json\n{"items": []}\n```') == {"items": []}
    cut = '{"items": [{"type": "step", "text": "a"}, {"type": "step", "text": "b"}, {"type": "st'
    assert len(nt.extract_json(cut)["items"]) == 2


def test_keys(monkeypatch):
    a, b = "AIza" + "A" * 35, "AQ.Ab8RN6KLKbQvyy6JzJKO8GVLmBqNs7rWihkebGjYO7wa"
    monkeypatch.setenv("GEMINI_API_KEY", f"{a}, {b} ; bad")
    monkeypatch.setenv("NOTES_API_KEY", a)
    assert nt.get_keys() == [a, b]


# ------------------------------------------------------------------ pausing on the writing (synthetic video)

def _fake_lecture(path: Path) -> None:
    """A 'notebook' video: page 1 is written line by line with a moving hand, then the page is turned
    and page 2 gets written; the hand rests on page 2 for a while (must NOT count as a new page)."""
    import cv2
    import numpy as np
    w, h, fps = 360, 640, 10
    out = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"mp4v"), fps, (w, h))
    blank = np.full((h, w, 3), 235, np.uint8)
    blank[:, :40] = blank[:, -40:] = 25          # dark side bars like a phone recording

    def page(lines, label):
        img = blank.copy()
        for k in range(lines):
            cv2.putText(img, f"{label} line {k} = {k * 7}", (55, 80 + 45 * k), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (120, 40, 20), 2)
        return img

    def hand(img, x, y):
        img = img.copy()
        cv2.ellipse(img, (x, y), (60, 90), 20, 0, 360, (60, 80, 110), -1)
        return img

    for k in range(1, 9):                         # page 1, 8 lines, hand moving between lines
        for f in range(fps * 2):
            out.write(hand(page(k, "A"), 200 + (f % 5) * 6, 120 + 45 * k) if f < fps else page(k, "A"))
    for f in range(fps * 2):                      # page being turned: big motion
        out.write(np.roll(page(8, "A"), f * 30, axis=1))
    for k in range(1, 6):                         # page 2
        for f in range(fps * 2):
            out.write(page(k, "B"))
    for f in range(fps * 5):                      # hand resting on page 2 (covers most lines)
        out.write(hand(page(5, "B"), 200, 200))
    for f in range(fps * 2):
        out.write(page(6, "B"))
    out.release()


def test_board_snapshots_finds_each_full_page(tmp_path):
    video = tmp_path / "lecture.mp4"
    _fake_lecture(video)
    shots = nt.board_snapshots(video)
    assert len(shots) == 2, [round(s["t"], 1) for s in shots]
    assert shots[0]["t"] >= 14            # page 1 taken when all 8 lines were there
    assert shots[1]["t"] >= 33            # page 2 taken at the end, with line 6 (after the hand left)
    for s in shots:
        assert s["jpeg"][:2] == b"\xff\xd8" and max(s["size"]) >= 1000   # enhanced and upscaled JPEG


def test_enhance_crops_to_writing():
    import numpy as np
    frame = np.full((360, 640, 3), 30, np.uint8)
    frame[100:200, 300:400] = 230
    mask = np.zeros((360, 640), bool)
    mask[120:180, 320:380] = True
    jpeg, size = nt.enhance(frame, mask)
    assert jpeg[:2] == b"\xff\xd8" and size[0] < 1000 and size[1] < 1000   # cropped, then scaled up 4x max


# ------------------------------------------------------------------ model calls (fake HTTP)

class Resp:
    def __init__(self, code, body):
        self.status_code, self._body, self.text, self.headers = code, body, json.dumps(body), {}

    def json(self):
        return self._body


def ok(items, title="T"):
    return Resp(200, {"candidates": [{"content": {"parts": [{"text": "thinking", "thought": True},
                                                            {"text": json.dumps({"title": title, "items": items})}]}}]})


def fake_post(monkeypatch, responses, seen=None):
    queue = list(responses)

    def post(url, headers=None, json=None, timeout=None, **_):
        if seen is not None:
            seen.append((url, headers.get("x-goog-api-key") if headers else None, json))
        return queue.pop(0)
    monkeypatch.setattr(nt.requests, "post", post)


def test_generate_retries_minute_limit_then_succeeds(monkeypatch):
    seen = []
    fake_post(monkeypatch, [Resp(429, {"error": {"message": "rate", "details": [{"retryDelay": "3s"}]}}),
                            Resp(503, {"error": {"message": "busy"}}), ok([{"type": "text", "text": "hi"}])], seen)
    assert nt.generate([{"text": "x"}], nt.NOTES_SCHEMA, ["AIzaK"])["items"][0]["text"] == "hi"
    assert len(seen) == 3 and f"{nt.DEFAULT_MODELS[0]}:" in seen[0][0]


def test_generate_daily_quota_moves_to_next_key_then_model(monkeypatch):
    seen = []
    daily = Resp(429, {"error": {"message": "Quota exceeded for metric generate_content_free_tier_requests, PerDay"}})
    fake_post(monkeypatch, [daily, daily, Resp(404, {"error": {"message": "model not found"}}),
                            Resp(404, {"error": {"message": "model not found"}}), ok([])], seen)
    assert nt.generate([{"text": "x"}], nt.NOTES_SCHEMA, ["AIza1", "AIza2"])["items"] == []
    assert [s[1] for s in seen[:2]] == ["AIza1", "AIza2"]
    assert f"{nt.DEFAULT_MODELS[1]}:" in seen[2][0] and f"{nt.DEFAULT_MODELS[2]}:" in seen[4][0]


def test_generate_minute_limit_hands_over_to_next_key(monkeypatch):
    seen = []
    fake_post(monkeypatch, [Resp(429, {"error": {"message": "rate", "details": [{"retryDelay": "30s"}]}}),
                            ok([{"type": "text", "text": "hi"}])], seen)
    assert nt.generate([{"text": "x"}], nt.NOTES_SCHEMA, ["AIza1", "AIza2"])["items"][0]["text"] == "hi"
    assert [s[1] for s in seen] == ["AIza1", "AIza2"]


def test_backup_reader_takes_over_when_every_gemini_key_is_used_up(monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "gsk_" + "a" * 30 + ",gsk_" + "b" * 30)
    daily = Resp(429, {"error": {"message": "limit PerDay"}})
    groq_full = Resp(429, {"error": {"message": "rate limit"}})
    groq_ok = lambda items: Resp(200, {"choices": [{"message": {"content": "<think>hmm {x}</think>" + json.dumps(
        {"title": "Backup", "items": items})}}]})
    seen = []
    fake_post(monkeypatch, [daily] * len(nt.DEFAULT_MODELS) + [groq_full,
              groq_ok([{"type": "question", "text": "Q1. Find p"}]), groq_ok([{"type": "step", "text": "p = 4"}])], seen)
    parts = []
    for k in range(6):                                   # 6 pages → 2 backup calls of 4 + 2 images
        parts += [{"text": f"Image {k + 1}"}, {"inlineData": {"mimeType": "image/jpeg", "data": "AAAA"}}]
    parts.append({"text": "Copy the writing."})
    out = nt.generate(parts, nt.NOTES_SCHEMA, ["AIza1"])
    groq = [s for s in seen if "groq.com" in s[0]]
    assert len(groq) == 3                                # first key full → second key, then the next chunk
    assert out["title"] == "Backup" and [i["text"] for i in out["items"]] == ["Find p", "p = 4"]
    assert sum(1 for c in groq[1][2]["messages"][0]["content"] if c["type"] == "image_url") == 4
    assert "JSON" in groq[1][2]["messages"][0]["content"][0]["text"]


def test_order_is_gemini_then_groq_then_openrouter(monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "gsk_" + "a" * 30)
    monkeypatch.setenv("OPENROUTER_API_KEY", "sk-or-v1-" + "c" * 40)
    monkeypatch.setattr(nt, "BACKUPS", [dict(nt.BACKUPS[0]), dict(nt.BACKUPS[1], models=["vision-model:free"])])
    daily = Resp(429, {"error": {"message": "limit PerDay"}})
    dead_key = Resp(401, {"error": {"message": "invalid api key"}})
    or_ok = Resp(200, {"choices": [{"message": {"content": json.dumps(
        {"title": "From OpenRouter", "items": [{"type": "answer", "text": "x = 4"}]})}}]})
    seen = []
    fake_post(monkeypatch, [daily] * len(nt.DEFAULT_MODELS) + [dead_key, or_ok], seen)
    out = nt.generate([{"text": "Image 1"}, {"inlineData": {"mimeType": "image/jpeg", "data": "AAAA"}},
                       {"text": "Copy the writing."}], nt.NOTES_SCHEMA, ["AIza1"])
    hosts = [("googleapis" if "googleapis" in s[0] else "groq" if "groq.com" in s[0] else "openrouter") for s in seen]
    assert hosts == ["googleapis"] * len(nt.DEFAULT_MODELS) + ["groq", "openrouter"]
    assert out["title"] == "From OpenRouter" and out["items"] == [{"type": "answer", "text": "x = 4"}]


def test_youtube_reading_never_goes_to_backup(monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "gsk_" + "a" * 30)
    fake_post(monkeypatch, [Resp(429, {"error": {"message": "limit PerDay"}})] * len(nt.DEFAULT_MODELS))
    with pytest.raises(nt.QuotaError):
        nt.generate([{"fileData": {"fileUri": "https://youtu.be/x"}}, {"text": "x"}], nt.NOTES_SCHEMA, ["AIza1"])


def test_generate_all_quota_gone(monkeypatch):
    fake_post(monkeypatch, [Resp(429, {"error": {"message": "limit PerDay"}})] * len(nt.DEFAULT_MODELS))
    with pytest.raises(nt.QuotaError):
        nt.generate([{"text": "x"}], nt.NOTES_SCHEMA, ["AIza1"])


def test_generate_video_access_error(monkeypatch):
    fake_post(monkeypatch, [Resp(400, {"error": {"message": "Cannot fetch content from the provided YouTube URL"}})])
    with pytest.raises(nt.VideoAccessError):
        nt.generate([{"text": "x"}], nt.NOTES_SCHEMA, ["AIza1"])


# ------------------------------------------------------------------ whole pipeline (fake model)

INFO = {"id": "abcdefghijk", "title": "Lec 4", "duration": 1300, "max_height": 360, "is_live": False}


def _shots(n, extra_every=0):
    return [{"t": i * 60, "ink": .02, "jpeg": b"\xff\xd8img%d" % i, "size": (1200, 1400),
             "extra": (b"\xff\xd8alt%d" % i) if extra_every and i % extra_every == 0 else None} for i in range(n)]


def test_snapshots_sent_as_images_in_groups_then_joined(monkeypatch):
    calls = []

    def gen(parts, schema, keys, **_):
        calls.append(parts)
        if "RAW NOTES" in parts[-1]["text"]:
            return {"title": "Joined", "items": json.loads(parts[0]["text"].split("RAW NOTES (JSON):\n", 1)[1])}
        k = len(calls)
        return {"title": "Revenue", "items": [{"type": "question", "text": f"Q from group {k}"},
                                               {"type": "step", "text": f"step {k}"}]}
    monkeypatch.setattr(nt, "generate", gen)
    out = nt.notes_from_snapshots(INFO, _shots(14, extra_every=2), ["AIzaK"], lambda *a: None)
    images = [p for p in calls[0] if "inlineData" in p]
    assert len(images) == nt.IMAGES_PER_REQUEST and images[0]["inlineData"]["mimeType"] == "image/jpeg"
    assert len(calls) == 3                       # 21 images → 2 groups + 1 join
    assert "ONLY WHAT IS WRITTEN" in calls[0][-1]["text"] and "Never solve" in calls[0][-1]["text"]
    assert out["title"] == "Joined" and [i["text"] for i in out["items"]][0] == "Q from group 1"


def test_make_notes_uses_snapshots(monkeypatch, tmp_path):
    monkeypatch.setattr(nt, "get_video_info", lambda vid: dict(INFO, id=vid))
    monkeypatch.setattr(nt, "download_video", lambda vid, folder, dur, progress=None: tmp_path / "v.mp4")
    monkeypatch.setattr(nt, "board_snapshots", lambda path, progress=None: _shots(3))
    monkeypatch.setattr(nt, "generate", lambda *a, **k: {"title": "Revenue", "items": [{"type": "question", "text": "Find p"}]})
    out = nt.make_notes("https://youtu.be/AAAAAAAAAAA", keys=["AIzaK"])
    assert out["title"] == "Revenue" and out["pages"] == 3 and out["video_id"] == "AAAAAAAAAAA"
    # second time comes from the cache, without any model call
    monkeypatch.setattr(nt, "generate", lambda *a, **k: pytest.fail("should be cached"))
    assert nt.make_notes("AAAAAAAAAAA", keys=["AIzaK"])["title"] == "Revenue"


def test_make_notes_falls_back_to_youtube_when_download_blocked(monkeypatch):
    seen = []
    monkeypatch.setattr(nt, "get_video_info", lambda vid: dict(INFO, id=vid, duration=700))

    def blocked(*a, **k):
        raise RuntimeError("Sign in to confirm you're not a bot")
    monkeypatch.setattr(nt, "download_video", blocked)

    def gen(parts, schema, keys, **_):
        seen.append(parts)
        return {"title": "T", "items": [{"type": "step", "text": f"s{len(seen)}"}]}
    monkeypatch.setattr(nt, "generate", gen)
    out = nt.make_notes("BBBBBBBBBBB", keys=["AIzaK"])
    assert seen[0][0]["fileData"]["fileUri"].endswith("BBBBBBBBBBB")
    assert seen[0][0]["videoMetadata"] == {"startOffset": "0s", "endOffset": "620s", "fps": 0.5}
    assert out["items"]


def test_make_notes_messages(monkeypatch):
    with pytest.raises(nt.NotesError, match="YouTube video link"):
        nt.make_notes("https://example.com", keys=["AIzaK"])
    with pytest.raises(nt.NotesError, match="not switched on"):
        nt.make_notes("CCCCCCCCCCC", keys=[])
    monkeypatch.setattr(nt, "get_video_info", lambda vid: dict(INFO, id=vid))
    monkeypatch.setattr(nt, "download_video", lambda *a, **k: Path("x"))
    monkeypatch.setattr(nt, "board_snapshots", lambda *a, **k: _shots(1))
    monkeypatch.setattr(nt, "generate", lambda *a, **k: (_ for _ in ()).throw(nt.QuotaError("x")))
    with pytest.raises(nt.NotesError, match="tomorrow"):
        nt.make_notes("DDDDDDDDDDD", keys=["AIzaK"])


def test_lectures_up_to_two_hours(monkeypatch):
    monkeypatch.setattr(nt, "get_video_info", lambda vid: dict(INFO, id=vid, duration=2 * 3600 + 60))
    with pytest.raises(nt.NotesError, match="up to 2 hours"):
        nt.make_notes("FFFFFFFFFFF", keys=["AIzaK"])


def _wait_for(cond, timeout=5.0):
    import time
    end = time.time() + timeout
    while not cond():
        assert time.time() < end, "timed out"
        time.sleep(0.01)


def test_same_lecture_is_made_once_for_every_student(monkeypatch):
    import threading
    gate, calls = threading.Event(), []

    def fake(link, keys, status):
        calls.append(link)
        status("Reading the writing", 0.6)
        gate.wait(5)
        return dict(SAMPLE, title=f"notes of {nt.parse_video_id(link)}")
    monkeypatch.setattr(nt, "make_notes", fake)
    students = [nt.Job(links=["https://youtu.be/GGGGGGGGGGG", f"https://youtu.be/HHHHHHHHHH{k}"]) for k in range(3)]
    for job in students:
        nt.start_job(job, [])
    _wait_for(lambda: all(job.fractions[0] == 0.6 for job in students))   # all see the one running lecture
    gate.set()
    _wait_for(lambda: all(job.done for job in students))
    assert sum(1 for c in calls if "GGGGGGGGGGG" in c) == 1
    assert all(job.results[0]["title"] == "notes of GGGGGGGGGGG" for job in students)
    assert [job.results[1]["title"][-1] for job in students] == ["0", "1", "2"]


def test_students_take_fair_turns_and_see_their_place(monkeypatch):
    import threading
    gate = threading.Event()

    def fake(link, keys, status):
        if "blk" in link:
            gate.wait(5)
        return dict(SAMPLE, title=link)
    monkeypatch.setattr(nt, "make_notes", fake)
    nt.start_job(nt.Job(links=[]), [])                        # make sure the workers are running
    blockers = nt.Job(links=[f"blk{k}" for k in range(len(nt._workers))])
    nt.start_job(blockers, [])
    _wait_for(lambda: blockers.id not in nt._pending)          # every worker is busy now
    big, small = nt.Job(links=["big1", "big2", "big3"]), nt.Job(links=["small1", "small2"])
    nt.start_job(big, [])
    nt.start_job(small, [])
    places = lambda job: [int(s.rsplit(" ", 1)[1]) for s in job.stages]
    assert places(big) == [1, 3, 5] and places(small) == [2, 4]   # one each, round and round
    gate.set()
    _wait_for(lambda: big.done and small.done and blockers.done)
    assert [r["title"] for r in small.results] == ["small1", "small2"]


def test_job_runs_links_in_parallel_and_keeps_order(monkeypatch):
    import threading
    running, peak = [0], [0]
    lock = threading.Lock()

    def fake(link, keys, status):
        with lock:
            running[0] += 1
            peak[0] = max(peak[0], running[0])
        import time
        time.sleep(0.05)
        status("Reading", 0.5)
        with lock:
            running[0] -= 1
        if "bad" in link:
            raise nt.NotesError("nope")
        if "boom" in link:
            raise RuntimeError("x")
        return dict(SAMPLE, title=link)
    monkeypatch.setattr(nt, "make_notes", fake)
    job = nt.Job(links=["a", "bad", "c", "boom", "e", "f", "g"])
    nt.run_job(job, [])
    assert job.done and job.progress == 1.0 and peak[0] > 1
    assert [r["title"] for r in job.results if r] == ["a", "c", "e", "f", "g"]
    assert [e[0] for e in job.errors] == [2, 4]
    both = nt.combine(job.results)
    assert both["lectures"] == 5 and both["items"][0] == {"type": "lecture", "text": "Lecture 1 · a"}


# ------------------------------------------------------------------ exports

def test_grouping():
    qs = [b for k, b in nt.grouped(SAMPLE["items"]) if k == "question"]
    assert [q["n"] for q in qs] == [1, 2]
    assert len(qs[0]["steps"]) == 3 and qs[0]["answers"][0]["text"] == "p = 40 cents"


def test_text_export():
    txt = nt.to_text(SAMPLE)
    assert "Q1. Problem, page 778" in txt and "Answer: p = 40 cents" in txt and "http" not in txt


def test_pdf_export():
    data = nt.to_pdf(SAMPLE)
    assert data[:5] == b"%PDF-" and len(data) > 2000
    many = nt.combine([dict(SAMPLE, title=f"L{i}") for i in range(8)])
    assert nt.to_pdf(many)[:5] == b"%PDF-"


def test_pdf_credit_on_every_page():
    pypdf = pytest.importorskip("pypdf")
    many = nt.combine([dict(SAMPLE, title=f"L{i}") for i in range(8)])
    pages = pypdf.PdfReader(io.BytesIO(nt.to_pdf(many))).pages
    assert len(pages) > 1 and all("Manufacturer: Muhammad Saif Ur Rehman" in p.extract_text() for p in pages)


def test_docx_export():
    from docx import Document
    doc = Document(io.BytesIO(nt.to_docx(SAMPLE)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Q1.  Problem, page 778" in text and "R = 10,000p − 125p²" in text
    assert all(s.footer.paragraphs[0].text == nt.CREDIT for s in doc.sections)       # every page, bottom right


def test_filename():
    assert nt.safe_filename('Lec: 4 / "Roots"?') == "Lec 4 Roots"


# ------------------------------------------------------------------ app screens

def _app():
    from streamlit.testing.v1 import AppTest
    return AppTest.from_file(str(HERE / "app.py"), default_timeout=30)


def _page(at):
    return " ".join(h.proto.body for h in at.get("html"))


def test_app_home_splash_and_about():
    at = _app().run()
    assert not at.exception
    assert "Turn lectures into <em>clean notes</em>" in _page(at)
    assert 'class="ink-splash"' in _page(at) and "Muhammad Saif Ur Rehman" in _page(at)   # start-up only
    at.run()
    assert 'class="ink-splash"' not in _page(at) and "Student · Pakistan" not in _page(at)
    assert "Developer <b>Muhammad Saif Ur Rehman</b>" in _page(at)                      # under the logo
    assert not at.get("popover") and len(at.sidebar.button) == 2                          # no options, no history
    next(b for b in at.sidebar.button if b.label == "About").click().run()
    assert "saifsaab317@gmail.com" in _page(at) and "Student · Pakistan" in _page(at)


def test_app_result_view():
    at = _app()
    at.session_state["splashed"] = True
    job = nt.Job(links=["https://youtu.be/AAAAAAAAAAA", "https://youtu.be/BBBBBBBBBBB"])
    job.results = [dict(SAMPLE), None]
    job.errors = [(2, job.links[1], "This video is private, removed or not available.")]
    job.fractions, job.done = [1.0, 1.0], True
    nt.JOBS[job.id] = job
    at.session_state["job_id"] = job.id
    at.run()
    assert not at.exception
    assert at.header[0].value == SAMPLE["title"]
    assert "Link 2" in at.error[0].value
    page = _page(at)
    assert "Q1." in page and "p = 40 cents" in page and "2 questions" in page
    assert [b.proto.label for b in at.get("download_button")] == ["PDF", "Word", "Text"]


def test_app_progress_offers_what_is_ready():
    at = _app()
    at.session_state["splashed"] = True
    job = nt.Job(links=["https://youtu.be/AAAAAAAAAAA", "https://youtu.be/BBBBBBBBBBB"])
    job.results = [dict(SAMPLE), None]
    job.fractions, job.stages = [1.0, 0.0], ["Done", "Waiting in line · number 3"]
    nt.JOBS[job.id] = job
    at.session_state["job_id"] = job.id
    at.run()
    assert not at.exception
    assert "number 3" in _page(at) and any("Ready now: 1 lecture" in c.value for c in at.caption)
    assert [b.proto.label for b in at.get("download_button")] == ["PDF", "Word", "Text"]


def test_app_bad_link_message():
    at = _app()
    at.session_state["splashed"] = True
    at.run()
    at.chat_input[0].set_value("hello sir").run()
    assert not at.exception
    assert "YouTube video link" in at.error[0].value
