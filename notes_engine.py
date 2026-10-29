"""Youtube Lectures Links To Notes engine: YouTube lecture links -> the teacher's writing as clean notes (PDF / Word / text).

Per video
  1. Download the picture only (best quality between 240p and 720p, picked automatically).
  2. "Pause" the video by itself: find every moment a page / board / screen is fullest, just before it
     is turned, wiped or scrolled, and keep the sharpest still of it (plus a second view when the
     teacher's hand hides part of it). Zoom to the writing, upscale, add contrast and sharpen.
  3. A free vision model (Google Gemini) reads those stills and writes out the question statements and
     the solutions exactly as written, nothing the teacher only says.
  If YouTube does not allow the download (it happens on some servers), the model watches the video
  straight from YouTube instead.

Nothing here imports Streamlit, so it is easy to test.
"""
from __future__ import annotations

import base64
import hashlib
import html
import io
import json
import os
import re
import shutil
import tempfile
import threading
import time
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import requests

APP_NAME = "Youtube Lectures Links To Notes"
VERSION = "2.0"
DEVELOPER = "Muhammad Saif Ur Rehman"
CREDIT = f"Manufacturer: {DEVELOPER} · Student · Pakistan"   # bottom-right of every PDF / Word page
HERE = Path(__file__).resolve().parent
FONTS = HERE / "fonts"
CACHE_VERSION = "v2"

API = "https://generativelanguage.googleapis.com"
DEFAULT_MODELS = ["gemini-3.6-flash", "gemini-3.8-flash", "gemini-flash-latest", "gemini-3.7-flash",
                  "gemini-3.5-flash", "gemini-2.5-flash", "gemini-flash-lite-latest"]
ITEM_TYPES = ["heading", "text", "definition", "formula", "question", "step", "answer"]

MAX_LINKS = 50
MAX_SECONDS = 2 * 3600    # longest lecture accepted
WORKERS = 2               # lectures worked on at the same time for ALL users (env NOTES_WORKERS overrides)
JOB_KEEP_SECONDS = 12 * 3600   # finished jobs stay reachable by their link this long
IMAGES_PER_REQUEST = 16
SEGMENT_SECONDS = 600     # direct-from-YouTube fallback: one request per 10 minutes
OVERLAP_SECONDS = 20
MIN_GAP_SECONDS = 6.5     # stay under the free tier's requests-per-minute limit


class NotesError(Exception):
    """A problem the student should read (shown as-is in the app)."""


class VideoAccessError(Exception):
    """The model could not open the YouTube link directly."""


class QuotaError(Exception):
    """Every key/model is out of free quota for today."""


# --------------------------------------------------------------------------- links

_ID = r"([A-Za-z0-9_-]{11})"
_LINK_PATTERNS = [
    rf"(?:v=|vi=){_ID}",
    rf"youtu\.be/{_ID}",
    rf"/(?:shorts|live|embed|v|e)/{_ID}",
]


def parse_video_id(link: str) -> str | None:
    link = (link or "").strip()
    if re.fullmatch(_ID, link):
        return link
    if "youtu" not in link.lower():
        return None
    for pat in _LINK_PATTERNS:
        m = re.search(pat, link)
        if m:
            return m.group(1)
    return None


def parse_playlist_id(link: str) -> str | None:
    m = re.search(r"[?&]list=([A-Za-z0-9_-]{10,})", link or "")
    return m.group(1) if m else None


def split_links(text: str) -> list[str]:
    """Every YouTube link in a pasted block (spaces, commas, numbers or new lines around them)."""
    found = re.findall(r"(?:https?://)?(?:[\w-]+\.)?(?:youtube\.com|youtu\.be|youtube-nocookie\.com)/\S+", text or "")
    if not found:
        found = [t for t in re.split(r"[\s,]+", text or "") if re.fullmatch(_ID, t)]
    out = []
    for f in found:
        f = f.rstrip(").,;'\"")
        if f not in out:
            out.append(f)
    return out


def watch_url(video_id: str) -> str:
    return f"https://www.youtube.com/watch?v={video_id}"


def expand_playlist(list_id: str, limit: int = MAX_LINKS) -> list[str]:
    try:
        import yt_dlp
        opts = {"quiet": True, "no_warnings": True, "extract_flat": True, "playlistend": limit, "skip_download": True}
        with yt_dlp.YoutubeDL(opts) as ydl:
            info = ydl.extract_info(f"https://www.youtube.com/playlist?list={list_id}", download=False)
        return [e["id"] for e in (info or {}).get("entries") or [] if e and e.get("id")][:limit]
    except Exception:
        return []


def fmt_time(seconds: float | int | None) -> str:
    if not seconds:
        return "—"
    s = int(seconds)
    h, m, s = s // 3600, s % 3600 // 60, s % 60
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


# --------------------------------------------------------------------------- video details + download

def get_video_info(video_id: str) -> dict:
    """Title, length and best quality. Never raises for a public video."""
    try:
        import yt_dlp
        opts = {"quiet": True, "no_warnings": True, "skip_download": True, "noplaylist": True}
        with yt_dlp.YoutubeDL(opts) as ydl:
            info = ydl.extract_info(watch_url(video_id), download=False)
        heights = [f.get("height") for f in info.get("formats") or [] if f.get("height") and f.get("vcodec") != "none"]
        return {"id": video_id, "title": info.get("title") or "Lecture", "duration": int(info.get("duration") or 0),
                "max_height": max(heights) if heights else 0, "is_live": bool(info.get("is_live"))}
    except Exception as e:
        if any(w in str(e).lower() for w in ("private video", "video unavailable", "removed", "does not exist")):
            raise NotesError("This video is private, removed or not available.")
    # Fallback: oEmbed for the title, the watch page for the length.
    out = {"id": video_id, "title": "Lecture", "duration": 0, "max_height": 0, "is_live": False}
    try:
        r = requests.get("https://www.youtube.com/oembed", params={"url": watch_url(video_id), "format": "json"}, timeout=15)
        if r.status_code == 200:
            out["title"] = r.json().get("title") or out["title"]
        elif r.status_code in (401, 403, 404):
            raise NotesError("This video is private, removed or not available.")
        page = requests.get(watch_url(video_id), headers={"Accept-Language": "en"}, timeout=20).text
        m = re.search(r'"lengthSeconds":"(\d+)"', page)
        if m:
            out["duration"] = int(m.group(1))
    except requests.RequestException:
        pass
    return out


def download_video(video_id: str, folder: Path, duration: int, progress=None) -> Path:
    """The picture only (no sound needed): best quality between 240p and 720p, or 480p for very long lectures."""
    import yt_dlp
    cap = 480 if duration > 90 * 60 else 720
    fmt = (f"bv*[height<={cap}][height>=240][vcodec^=avc1]/bv*[height<={cap}][height>=240][vcodec!^=av01]/"
           f"bv*[height<={cap}][height>=240]/bv*[height<={cap}]/b[height<={cap}]/b")

    def hook(d):
        if progress and d.get("status") == "downloading":
            total = d.get("total_bytes") or d.get("total_bytes_estimate") or 0
            if total:
                progress(d.get("downloaded_bytes", 0) / total)

    opts = {"quiet": True, "no_warnings": True, "format": fmt, "outtmpl": str(folder / "%(id)s.%(ext)s"),
            "noplaylist": True, "progress_hooks": [hook], "retries": 5, "fragment_retries": 5,
            "concurrent_fragment_downloads": 4}
    with yt_dlp.YoutubeDL(opts) as ydl:
        ydl.download([watch_url(video_id)])
    files = [p for p in folder.glob(f"{video_id}.*") if not p.name.endswith(".part")]
    if not files:
        raise NotesError("Could not download this video.")
    return max(files, key=lambda p: p.stat().st_size)


# --------------------------------------------------------------------------- pausing on the writing

def _ink_mask(gray):
    """(ink, surface) boolean masks.

    Ink is only looked for on the writing surface: on paper / whiteboard (dark strokes on a bright
    area), or on a blackboard (bright strokes on a dark area) when there is hardly any bright area.
    Dark side bars of phone recordings, the desk, the room and the teacher's hand are not surface.
    """
    import cv2
    import numpy as np
    g = gray.astype(np.int16)
    local = cv2.GaussianBlur(gray, (0, 0), 6).astype(np.int16)
    paper = local > 115
    light = paper.mean() >= 0.12
    if light:
        surface, mask = paper, ((local - g) > 22) & paper
        level = float(np.percentile(gray[paper], 90)) if paper.any() else 200.0
        dark = (g < level - 45).astype(np.uint8)
    else:
        surface, mask = local < 115, ((g - local) > 22) & (local < 115)
        dark = np.zeros_like(gray, np.uint8)
    mask = cv2.morphologyEx(mask.astype(np.uint8), cv2.MORPH_OPEN, np.ones((2, 2), np.uint8)) > 0
    surface = cv2.dilate(surface.astype(np.uint8), np.ones((5, 5), np.uint8)) > 0   # ink sits on its edge too
    # Pen strokes are thin; anything thick and dark (hand, sleeve, pen body) is not writing.
    blobs = cv2.morphologyEx(dark, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (11, 11)))
    blobs = cv2.dilate(blobs, np.ones((7, 7), np.uint8)) > 0
    return mask & ~blobs, surface & ~blobs


def enhance(frame, mask=None):
    """Make blurry writing easier to read: zoom to the writing, upscale, contrast, sharpen. -> (jpeg, size)"""
    import cv2
    import numpy as np
    h, w = frame.shape[:2]
    if mask is not None and mask.any():
        ys, xs = np.nonzero(mask)
        sy, sx = h / mask.shape[0], w / mask.shape[1]
        y0, y1 = np.percentile(ys, [0.5, 99.5]) * sy
        x0, x1 = np.percentile(xs, [0.5, 99.5]) * sx
        py, px = 0.05 * h, 0.05 * w
        y0, y1 = int(max(0, y0 - py)), int(min(h, y1 + py))
        x0, x1 = int(max(0, x0 - px)), int(min(w, x1 + px))
        if (y1 - y0) * (x1 - x0) < 0.7 * h * w and y1 - y0 > 40 and x1 - x0 > 40:
            frame = frame[y0:y1, x0:x1]
    h, w = frame.shape[:2]
    scale = min(4.0, 1600 / max(h, w))
    if scale > 1.05:
        frame = cv2.resize(frame, (int(w * scale), int(h * scale)), interpolation=cv2.INTER_CUBIC)
    lab = cv2.cvtColor(frame, cv2.COLOR_BGR2LAB)
    lab[:, :, 0] = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(lab[:, :, 0])
    frame = cv2.cvtColor(lab, cv2.COLOR_LAB2BGR)
    blur = cv2.GaussianBlur(frame, (0, 0), 1.4)
    frame = cv2.addWeighted(frame, 1.7, blur, -0.7, 0)
    ok, buf = cv2.imencode(".jpg", frame, [cv2.IMWRITE_JPEG_QUALITY, 90])
    return buf.tobytes(), frame.shape[1::-1]


def board_snapshots(path, progress=None, max_shots: int | None = None) -> list[dict]:
    """Stop the video by itself at every moment a page/board is about to be wiped, turned or
    scrolled, keep the fullest still of it (plus a second view where the hand is elsewhere),
    and enhance them for reading.
    """
    import cv2
    import numpy as np
    cap = cv2.VideoCapture(str(path))
    fps = cap.get(cv2.CAP_PROP_FPS) or 25
    total = cap.get(cv2.CAP_PROP_FRAME_COUNT) or 1
    every = max(1, round(fps * (1.0 if total / fps <= 40 * 60 else 2.0)))
    max_shots = max_shots or max(60, int(total / fps / 60))   # long lectures keep more pages (2 h → 120)
    kernel = np.ones((5, 5), np.uint8)
    shots: list[dict] = []
    best = alt = None
    prev_gray = None
    pending = idx = 0

    def close():
        if best is None or best["ink"] <= 0.004:
            return
        jpeg, size = enhance(best["frame"], best["mask"])
        shot = {"t": best["t"], "ink": best["ink"], "jpeg": jpeg, "size": size, "extra": None}
        # Second view: the least-covered still with nearly all the writing, if the hand is elsewhere.
        if alt is not None and abs(alt["t"] - best["t"]) > 3 and \
                float((best["covered"] ^ alt["covered"]).mean()) > 0.02:
            shot["extra"] = enhance(alt["frame"], best["mask"])[0]
        shots.append(shot)

    while cap.grab():
        idx += 1
        if (idx - 1) % every:
            continue
        ok, frame = cap.retrieve()
        if not ok:
            continue
        if progress and idx % (every * 20) == 1:
            progress(min(idx / total, 1.0))
        h, w = frame.shape[:2]
        gray = cv2.cvtColor(cv2.resize(frame, (480, max(1, int(480 * h / w)))), cv2.COLOR_BGR2GRAY)
        motion = float(cv2.absdiff(gray, prev_gray).mean()) if prev_gray is not None else 99.0
        prev_gray = gray
        if motion >= 8:
            continue                      # camera moving / page being turned: wait
        mask, surface = _ink_mask(gray)
        ink = float(mask.mean())
        if best is not None and motion < 3:   # only decide "page changed" on a still picture
            # Only judge writing that is not hidden right now (a hand resting on the page is not an erase).
            seen = best["mask"] & surface
            if seen.sum() >= 0.3 * best["mask"].sum():
                grown = cv2.dilate(mask.astype(np.uint8), kernel) > 0
                lost = float((seen & ~grown).sum()) / max(1, int(seen.sum()))
                pending = pending + 1 if lost >= 0.4 else 0
            if pending >= 3:              # the writing is really gone: page turned / board wiped / scrolled
                close()
                best, alt, pending = None, None, 0
        if pending:
            continue
        t = (idx - 1) / fps
        # Writing only grows until the page changes, so a later still with about as much ink wins.
        if best is None or ink >= best["ink"] * 0.97:
            best = {"t": t, "ink": ink, "frame": frame.copy(), "mask": mask, "covered": ~surface}
            alt = None
        elif ink >= 0.85 * best["ink"] and (alt is None or surface.mean() > alt["open"]):
            alt = {"t": t, "frame": frame.copy(), "open": float(surface.mean()), "covered": ~surface}
    close()
    cap.release()
    if len(shots) > max_shots:
        import numpy as np
        keep = np.linspace(0, len(shots) - 1, max_shots).round().astype(int)
        shots = [shots[i] for i in sorted(set(keep))]
    return shots


# --------------------------------------------------------------------------- text cleaning

_SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
_SUB = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")
_LATEX_WORDS = {
    r"\times": "×", r"\cdot": "·", r"\div": "÷", r"\pm": "±", r"\mp": "∓", r"\leq": "≤", r"\le": "≤",
    r"\geq": "≥", r"\ge": "≥", r"\neq": "≠", r"\ne": "≠", r"\approx": "≈", r"\infty": "∞", r"\pi": "π",
    r"\theta": "θ", r"\alpha": "α", r"\beta": "β", r"\gamma": "γ", r"\delta": "δ", r"\Delta": "Δ",
    r"\lambda": "λ", r"\mu": "μ", r"\sigma": "σ", r"\Sigma": "Σ", r"\sum": "Σ", r"\int": "∫",
    r"\rightarrow": "→", r"\to": "→", r"\Rightarrow": "⇒", r"\implies": "⇒", r"\therefore": "∴",
    r"\because": "∵", r"\degree": "°", r"\circ": "°", r"\in": "∈", r"\subset": "⊂", r"\cup": "∪",
    r"\cap": "∩", r"\angle": "∠", r"\triangle": "△", r"\perp": "⊥", r"\parallel": "∥", r"\phi": "φ",
    r"\omega": "ω", r"\partial": "∂", r"\nabla": "∇", r"\sqrt": "√", r"\left": "", r"\right": "",
    r"\,": " ", r"\;": " ", r"\quad": "  ", r"\!": "",
}
_URL = re.compile(r"(https?://\S+|www\.\S+|\b[A-Za-z][\w-]*\.(?:com|net|org|pk|io)(?:/\S*)?\b)", re.I)
_PROMO = re.compile(r"\b(subscribe|like (?:and|&) share|bell icon|follow (?:me|us)|whats ?app|join (?:my|our) channel|telegram|instagram|facebook)\b", re.I)
_SCRIPT_RUNS = re.compile(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF\u0900-\u097F]+")


def tidy_math(s: str) -> str:
    """Turn stray LaTeX into plain readable maths (x², √(x), (a)/(b))."""
    if not s:
        return ""
    s = s.replace("$$", "").replace("$", "").replace(r"\(", "").replace(r"\)", "").replace(r"\[", "").replace(r"\]", "")
    for _ in range(3):
        s = re.sub(r"\\[dt]?frac\{([^{}]*)\}\{([^{}]*)\}", lambda m: f"({m.group(1)})/({m.group(2)})", s)
        s = re.sub(r"\\sqrt\[(\d+)\]\{([^{}]*)\}", lambda m: {"3": "∛", "4": "∜"}.get(m.group(1), f"{m.group(1)}√") + f"({m.group(2)})", s)
        s = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\(?:text|mathrm|mathbf|operatorname)\{([^{}]*)\}", r"\1", s)
    for k in sorted(_LATEX_WORDS, key=len, reverse=True):
        s = re.sub(re.escape(k) + r"(?![A-Za-z])", lambda _m, v=_LATEX_WORDS[k]: v, s)
    s = re.sub(r"\^\{([0-9+\-=()n]+)\}", lambda m: m.group(1).translate(_SUP), s)
    s = re.sub(r"\^([0-9n])(?![0-9])", lambda m: m.group(1).translate(_SUP), s)
    s = re.sub(r"_\{([0-9]+)\}", lambda m: m.group(1).translate(_SUB), s)
    s = re.sub(r"(?<=[A-Za-z])_([0-9])(?![0-9])", lambda m: m.group(1).translate(_SUB), s)
    s = re.sub(r"\^\{([^{}]+)\}", r"^(\1)", s)
    s = re.sub(r"\(([A-Za-z0-9.]+)\)/\(([A-Za-z0-9.]+)\)", r"\1/\2", s)
    return re.sub(r"\\([A-Za-z]+)", r"\1", s)


def clean_text(s: str) -> str:
    s = html.unescape(str(s or ""))
    s = _URL.sub("", s)
    s = _SCRIPT_RUNS.sub("", s)
    s = tidy_math(s)
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r" *\n *", "\n", s).strip(" \n-–—:;,")
    return s


def clean_items(items: list) -> list[dict]:
    out = []
    for it in items or []:
        if not isinstance(it, dict):
            continue
        typ = str(it.get("type") or "text").lower().strip()
        typ = {"title": "heading", "subheading": "heading", "solution": "step", "working": "step",
               "final answer": "answer", "result": "answer", "rule": "formula"}.get(typ, typ)
        if typ not in ITEM_TYPES:
            typ = "text"
        text = clean_text(it.get("text", ""))
        if typ == "question":
            text = re.sub(r"^(?:(?:q|ques|question)\.?\s*(?:no\.?|#)?\s*\d+[a-z]?\s*[:.)-]?|question\s*[:.-])\s*", "",
                          text, flags=re.I) or text
        if typ == "answer":
            text = re.sub(r"^(?:final\s+)?ans(?:wer)?\s*[.:=-]\s*", "", text, flags=re.I) or text
        if typ == "step":
            text = re.sub(r"^(?:step\s*\d+\s*[:.)-]\s*)", "", text, flags=re.I) or text
        if not text or _PROMO.search(text) and len(text) < 160 or re.fullmatch(r"(?i)sol(ution)?\.?:?", text):
            continue
        out.append({"type": typ, "text": text})
    return out


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9=+\-*/^√π²³]", "", s.lower())


def drop_repeats(items: list[dict]) -> list[dict]:
    """Remove an item that exactly repeats the one before it (board still on screen)."""
    out: list[dict] = []
    for it in items:
        key = _norm(it["text"])
        if key and not (out and out[-1]["type"] == it["type"] and _norm(out[-1]["text"]) == key):
            out.append(it)
    return out


def join_parts(raw: list[dict], new: list[dict], look_back: int = 40) -> list[dict]:
    """Append the next part, dropping what it repeats from the overlap at the cut.

    Only the leading run of the new part is checked, so a step such as "x = 2" that
    really appears in two different questions is kept.
    """
    recent = {(o["type"], _norm(o["text"])) for o in raw[-look_back:]}
    k = 0
    while k < min(len(new), look_back) and (new[k]["type"], _norm(new[k]["text"])) in recent:
        k += 1
    return drop_repeats(raw + new[k:])


# --------------------------------------------------------------------------- model calls

def get_keys() -> list[str]:
    raw = ",".join(os.environ.get(k, "") for k in ("NOTES_API_KEY", "GEMINI_API_KEY", "GOOGLE_API_KEY"))
    keys = []
    for k in re.split(r"[\s,;]+", raw):   # "AIza…" (classic) and "AQ.…" (newer) Google keys
        if re.fullmatch(r"[A-Za-z0-9._-]{30,}", k) and k not in keys:
            keys.append(k)
    return keys


def get_models() -> list[str]:
    custom = [m.strip() for m in os.environ.get("NOTES_MODEL", "").split(",") if m.strip()]
    return custom + [m for m in DEFAULT_MODELS if m not in custom]


_last_call: dict[str, float] = {}
_call_lock = threading.Lock()
_dead: dict[tuple[str, str], float] = {}   # (key, model) -> time it ran out of daily quota


def _pace(key: str) -> None:
    try:   # a paid key can go faster: NOTES_MIN_GAP = "0.5" in the secrets
        gap = float(os.environ.get("NOTES_MIN_GAP") or MIN_GAP_SECONDS)
    except ValueError:
        gap = MIN_GAP_SECONDS
    with _call_lock:
        wait = _last_call.get(key, 0) + gap - time.time()
        _last_call[key] = time.time() + max(wait, 0)
    if wait > 0:
        time.sleep(wait)


def _retry_delay(err: dict) -> float:
    for d in err.get("details") or []:
        delay = d.get("retryDelay")
        if delay:
            try:
                return float(str(delay).rstrip("s"))
            except ValueError:
                pass
    return 20.0


def _is_daily(err: dict) -> bool:
    text = json.dumps(err).lower()
    return "perday" in text or "per day" in text or "daily" in text


def _gen_config(model: str, schema: dict, max_tokens: int) -> dict:
    big = any(t in model for t in ("2.5", "latest", "-3"))
    cfg = {
        "temperature": 0.1,
        "responseMimeType": "application/json",
        "responseSchema": schema,
        "maxOutputTokens": max_tokens if big else min(max_tokens, 8192),
    }
    if "2.5" in model and "lite" not in model:
        cfg["thinkingConfig"] = {"thinkingBudget": 8192}
    return cfg


def extract_json(text: str) -> dict:
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    start = text.find("{")
    if start < 0:
        raise ValueError("no JSON")
    body = text[start:]
    try:
        return json.loads(body[: body.rfind("}") + 1])
    except json.JSONDecodeError:
        pass
    # Cut-off answer: keep every complete item.
    cut = body.rfind("},")
    while cut > 0:
        try:
            return json.loads(body[: cut + 1] + "]}")
        except json.JSONDecodeError:
            cut = body.rfind("},", 0, cut)
    raise ValueError("bad JSON")


def generate(parts: list, schema: dict, keys: list[str], models: list[str] | None = None,
             max_tokens: int = 65536, log=None, timeout: int = 300) -> dict:
    """One JSON answer. Gemini first (every key, every model); when all of them are used up or busy,
    the backup reader (Groq) takes over — except for videos read straight from YouTube, which only Gemini can do."""
    try:
        return _gemini_generate(parts, schema, keys, models, max_tokens, log, timeout)
    except (QuotaError, NotesError) as first:
        backup = get_groq_keys()
        if not backup or any("fileData" in p for p in parts):
            raise
        try:
            return _groq_generate(parts, backup, log)
        except Exception:
            raise first


def _gemini_generate(parts: list, schema: dict, keys: list[str], models: list[str] | None = None,
                     max_tokens: int = 65536, log=None, timeout: int = 300) -> dict:
    """One JSON answer from the first key/model that works. Handles limits and retries."""
    models = models or get_models()
    if not keys:
        raise QuotaError("no key")
    last = "unknown error"
    for model in models:
        for key in keys:
            if time.time() - _dead.get((key, model), 0) < 3600:
                continue
            for attempt in range(6):
                _pace(key)
                body = {"contents": [{"role": "user", "parts": parts}],
                        "generationConfig": _gen_config(model, schema, max_tokens)}
                try:
                    r = requests.post(f"{API}/v1beta/models/{model}:generateContent",
                                      headers={"x-goog-api-key": key, "Content-Type": "application/json"},
                                      json=body, timeout=timeout)
                except requests.RequestException as e:
                    last = f"network: {e}"
                    if attempt >= 1:
                        break      # this model is hanging: try the next one
                    time.sleep(5)
                    continue
                if r.status_code == 200:
                    data = r.json()
                    cands = data.get("candidates") or []
                    if not cands:
                        last = f"blocked: {data.get('promptFeedback')}"
                        break
                    text = "".join(p.get("text", "") for p in (cands[0].get("content") or {}).get("parts") or []
                                   if not p.get("thought"))
                    try:
                        return extract_json(text)
                    except ValueError:
                        last = f"unreadable answer ({cands[0].get('finishReason')})"
                        if attempt >= 1:
                            break
                        continue
                try:
                    err = r.json().get("error") or {}
                except ValueError:
                    err = {"message": r.text[:300]}
                msg = str(err.get("message", ""))
                last = f"{r.status_code}: {msg[:300]}"
                if log:
                    log(f"{model}: {last}")
                if r.status_code == 429:
                    if _is_daily(err):
                        _dead[(key, model)] = time.time()
                        break
                    if key != keys[-1]:
                        break      # this key is full for the minute: the next key takes the request
                    time.sleep(min(_retry_delay(err) + 2, 65))
                    continue
                if r.status_code in (500, 502, 503, 504):
                    if attempt >= 1:
                        break      # busy ("high demand"): the next model is usually free
                    time.sleep(4)
                    continue
                if r.status_code == 404 or "not found" in msg.lower() and "model" in msg.lower():
                    break  # model not offered: next model
                if r.status_code in (401, 403) or "api key" in msg.lower():
                    break  # bad key: next key
                if r.status_code == 400:
                    low = msg.lower()
                    if any(w in low for w in ("youtube", "video", "file", "uri", "fetch", "private", "unavailable")):
                        raise VideoAccessError(msg)
                    if "thinking" in low or "responseschema" in low.replace("_", "") or "max_output" in low or "maxoutput" in low:
                        break  # this model doesn't accept the settings: next model
                    raise NotesError(f"The notes service refused this request ({msg[:200]}).")
                break
    if "429" in last or not any(time.time() - _dead.get((k, m), 0) >= 3600 for k in keys for m in models):
        raise QuotaError(last)
    raise NotesError(f"The notes service is not answering right now ({last[:200]}). Please try again in a few minutes.")


# --------------------------------------------------------------------------- backup reader (Groq)

GROQ_API = "https://api.groq.com/openai/v1/chat/completions"
GROQ_MODELS = ["qwen/qwen3.8-27b"]     # reads images; env GROQ_MODEL adds/overrides
GROQ_IMAGES_PER_REQUEST = 4
_JSON_SHAPE = ('\n\nReply with JSON only, exactly this shape: {"title": "...", "items": [{"type": '
               '"heading|text|definition|formula|question|step|answer", "text": "..."}]}')


def get_groq_keys() -> list[str]:
    keys = []
    for k in re.split(r"[\s,;]+", os.environ.get("GROQ_API_KEY", "")):
        if re.fullmatch(r"gsk_[A-Za-z0-9]{20,}", k) and k not in keys:
            keys.append(k)
    return keys


def get_groq_models() -> list[str]:
    custom = [m.strip() for m in os.environ.get("GROQ_MODEL", "").split(",") if m.strip()]
    return custom + [m for m in GROQ_MODELS if m not in custom]


def _groq_call(content: list, keys: list[str], log=None) -> dict:
    last = "unknown error"
    for model in get_groq_models():
        for key in keys:
            extra = {"response_format": {"type": "json_object"}, "reasoning_format": "hidden"}
            for attempt in range(4):
                body = {"model": model, "messages": [{"role": "user", "content": content}],
                        "temperature": 0.1, "max_completion_tokens": 8192} | extra
                try:
                    r = requests.post(GROQ_API, headers={"Authorization": f"Bearer {key}"}, json=body, timeout=180)
                except requests.RequestException as e:
                    last = f"network: {e}"
                    time.sleep(3)
                    continue
                if r.status_code == 200:
                    text = ((r.json().get("choices") or [{}])[0].get("message") or {}).get("content") or ""
                    try:
                        return extract_json(re.sub(r"<think>.*?</think>", "", text, flags=re.S))
                    except ValueError:
                        last = "unreadable answer"
                        continue
                last = f"{r.status_code}: {r.text[:200]}"
                if log:
                    log(f"groq {model}: {last}")
                low = r.text.lower()
                if r.status_code == 400 and any(p in low for p in extra):   # setting not accepted: drop it
                    extra = {k: v for k, v in extra.items() if k not in low}
                    continue
                if r.status_code == 429:
                    try:
                        wait = float(r.headers.get("retry-after") or 5)
                    except ValueError:
                        wait = 5.0
                    if key != keys[-1] or wait > 30:
                        break      # this key is full: the next key
                    time.sleep(wait)
                    continue
                if r.status_code in (500, 502, 503, 504):
                    time.sleep(3)
                    continue
                break              # bad key / model not offered: next one
    raise NotesError(f"The backup reader is not answering ({last[:150]}).")


def _groq_generate(parts: list, keys: list[str], log=None) -> dict:
    """The same request for Groq: a few images per call, joined in order."""
    images = [p["inlineData"] for p in parts if "inlineData" in p]
    prompt = "\n\n".join(p["text"] for p in parts if "text" in p and not re.fullmatch(r"Image \d+", p["text"]))
    prompt += _JSON_SHAPE
    if not images:
        return _groq_call([{"type": "text", "text": prompt}], keys, log)
    title, raw = "", []
    step = GROQ_IMAGES_PER_REQUEST
    for a in range(0, len(images), step):
        chunk = images[a:a + step]
        note = f"\n(These are images {a + 1}–{a + len(chunk)} of {len(images)}, in order.)" if len(images) > step else ""
        content = [{"type": "text", "text": prompt + note}] + [
            {"type": "image_url", "image_url": {"url": f"data:{im['mimeType']};base64,{im['data']}"}} for im in chunk]
        data = _groq_call(content, keys, log)
        title = title or clean_text(data.get("title") or "")
        raw = join_parts(raw, clean_items(data.get("items")))
    return {"title": title, "items": raw}


# --------------------------------------------------------------------------- prompts

ITEM_SCHEMA = {
    "type": "OBJECT",
    "properties": {"type": {"type": "STRING", "enum": ITEM_TYPES}, "text": {"type": "STRING"}},
    "required": ["type", "text"],
}
NOTES_SCHEMA = {
    "type": "OBJECT",
    "properties": {"title": {"type": "STRING"}, "items": {"type": "ARRAY", "items": ITEM_SCHEMA}},
    "required": ["title", "items"],
}

MATH_RULES = """Maths formatting (plain Unicode text, never LaTeX, never $ signs):
x², x³, xⁿ, x^(n+1) · √(x+1) · fractions (a+b)/(c−d) or 3/4 · × ÷ ± ≤ ≥ ≠ ≈ ∞ π θ Δ ∫ Σ → ⇒ ∴ ° ·
subscripts a₁, aₙ · keep the teacher's own letters, numbering like ① ② or (1) (2), arrows and units."""

WHAT_TO_WRITE = """WRITE ONLY WHAT IS WRITTEN — the teacher's talk is not wanted:
- every question / problem statement, complete and word for word, as ONE "question" item (put its parts
  (a), (b), (c) on new lines inside the same item; keep a book reference such as "Problem, page 778");
- every line of the solution as its own "step", in the order written, exactly as written (same numbers,
  signs, letters and equation labels); a final result line as "answer";
- a topic name written as a title → "heading"; a written formula on its own → "formula"; a written definition → "definition".
Never solve, add, explain, summarise or "improve" anything. If the writing stops mid-solution, stop there too.
Ignore: cover / title pages (student names, roll numbers, registration numbers, department, "submitted to"),
"Page # 1" style page labels, the teacher's name / signature / "folder" credit lines,
old writing from another page visible at the edge, the hand and pen,
video-call placeholders, logos, links, channel names.
The language of the notes is English, as written on the page."""


def snapshot_prompt(title: str, n_images: int, part: int, parts: int) -> str:
    return f"""These {n_images} images are stills taken automatically from a recorded lecture ("{title}"){f", group {part} of {parts}" if parts > 1 else ""}, in time order.
Each still was taken when a page / board / screen had the most writing on it, just before it was turned, wiped or scrolled.
Consecutive stills often show the SAME page with more writing added, and a page can appear twice (a second view where the
hand covers a different part) — write every line only ONCE, in the order the teacher wrote it.

The photos are low quality and the handwriting is blurry. Read like an expert who knows the subject:
look at every image of the same page, use the spelling of English words, the meaning of the problem and the maths itself
(each line follows from the one before it) to decide unclear words, digits and signs. Do not guess wildly: if a symbol
truly cannot be read, write [?] in its place.

{WHAT_TO_WRITE}
{MATH_RULES}
"title": a short, clear title for what is written (for example "Bus fare revenue — demand function").
Before answering, compare every line you wrote with the images once more: numbers, signs, powers, brackets."""


def segment_prompt(title: str, start: int, end: int, duration: int) -> str:
    return f"""You are copying a teacher's writing from part of a recorded lecture ("{title}"),
{fmt_time(start)} to {fmt_time(end)} of {fmt_time(duration) if duration else 'the video'}.
Watch every second of this part. Copy what is WRITTEN on the page / board / screen, in the order it is written.
The writing may be blurry: look at several frames of the same page and use the maths itself to read unclear symbols.
If something is crossed out or erased as a mistake, keep only the corrected version. If the same page stays on screen, write it once.
When writing continues from earlier in the lecture, just continue it.

{WHAT_TO_WRITE}
{MATH_RULES}
If this part has no writing, return {{"title": "", "items": []}}."""


def merge_prompt(title: str, raw_items: list[dict]) -> str:
    return f"""Below is a teacher's writing copied from consecutive, slightly overlapping parts of ONE lecture ("{title}"), in order.
Make the final notes:
- remove lines repeated because of the overlaps or because the same page was visible twice;
- join a question statement or a solution that was split across parts; keep lines in the original order;
- KEEP every distinct line. Do NOT summarise, shorten, solve, explain, reword maths or change numbers;
- give a short clear title.
{MATH_RULES}

RAW NOTES (JSON):
{json.dumps(raw_items, ensure_ascii=False)}"""


def polish(title: str, raw: list[dict], keys: list[str]) -> dict:
    """Join parts into one clean set of notes. The raw copy wins if anything got lost."""
    notes = {"title": title, "items": raw}
    try:
        merged = generate([{"text": merge_prompt(title, raw)}], NOTES_SCHEMA, keys)
    except (QuotaError, NotesError, VideoAccessError):
        return notes
    m_items = drop_repeats(clean_items(merged.get("items")))
    kept_chars = sum(len(x["text"]) for x in m_items) / max(1, sum(len(x["text"]) for x in raw))
    if m_items and (len(m_items) >= 0.8 * len(raw) or kept_chars >= 0.85):
        notes["items"] = m_items
    notes["title"] = clean_text(merged.get("title") or "") or title
    return notes


# --------------------------------------------------------------------------- reading the writing

def notes_from_snapshots(info: dict, shots: list[dict], keys: list[str], status) -> dict:
    """Send the stills (in groups) to the vision model and join what it reads."""
    images = []
    for s in shots:
        images.append(s["jpeg"])
        if s.get("extra"):
            images.append(s["extra"])
    groups = [images[i:i + IMAGES_PER_REQUEST] for i in range(0, len(images), IMAGES_PER_REQUEST)]
    title = info.get("title") or "Lecture"
    raw: list[dict] = []
    found_title = ""
    for g, imgs in enumerate(groups):
        status(f"Reading the writing{f' · part {g + 1} of {len(groups)}' if len(groups) > 1 else ''}",
               0.55 + 0.4 * g / len(groups))
        parts = []
        for k, jpeg in enumerate(imgs):
            parts.append({"text": f"Image {k + 1}"})
            parts.append({"inlineData": {"mimeType": "image/jpeg", "data": base64.b64encode(jpeg).decode()}})
        parts.append({"text": snapshot_prompt(title, len(imgs), g + 1, len(groups))})
        data = generate(parts, NOTES_SCHEMA, keys)
        found_title = found_title or clean_text(data.get("title") or "")
        raw = join_parts(raw, clean_items(data.get("items")))
    if len(groups) > 1 and raw:
        status("Joining the pages", 0.96)
        return polish(found_title or title, raw, keys)
    return {"title": found_title or title, "items": raw}


def notes_from_youtube(info: dict, keys: list[str], status) -> dict:
    """Fallback when the video can't be downloaded: the model watches it straight from YouTube."""
    vid, duration, title = info["id"], info.get("duration") or 0, info.get("title") or "Lecture"
    if not duration:
        segs = [(0, 0)]
    else:
        segs, start = [], 0
        while start < duration:
            end = min(start + SEGMENT_SECONDS, duration)
            if duration - end < 90:
                end = duration
            segs.append((start, min(end + OVERLAP_SECONDS, duration)))
            start = end
    fps = 0.5 if duration <= 3600 else 0.33
    raw: list[dict] = []
    found_title = ""
    i = 0
    while i < len(segs):
        s, e = segs[i]
        status(f"Reading the writing · part {i + 1} of {len(segs)}", 0.1 + 0.8 * i / len(segs))
        part = {"fileData": {"fileUri": watch_url(vid)}}
        meta = ({"startOffset": f"{s}s", "endOffset": f"{e}s"} if e else {}) | ({"fps": fps} if fps else {})
        if meta:
            part["videoMetadata"] = meta
        try:
            data = generate([part, {"text": segment_prompt(title, s, e, duration)}], NOTES_SCHEMA, keys)
        except NotesError:
            if not fps or raw:
                raise
            fps = None     # some models refuse a custom frame rate: use their default
            continue
        except VideoAccessError:
            raise NotesError("This video could not be opened.")
        found_title = found_title or clean_text(data.get("title") or "")
        raw = join_parts(raw, clean_items(data.get("items")))
        i += 1
    if len(segs) > 1 and raw:
        status("Joining the pages", 0.95)
        return polish(found_title or title, raw, keys)
    return {"title": found_title or title, "items": raw}


# --------------------------------------------------------------------------- cache + one video

def cache_dir() -> Path:
    d = Path(os.environ.get("NOTES_CACHE_DIR") or Path(tempfile.gettempdir()) / "lecture_to_notes_cache")
    d.mkdir(parents=True, exist_ok=True)
    return d


def _cache_path(video_id: str) -> Path:
    return cache_dir() / f"{video_id}-{CACHE_VERSION}.json"


def load_cached(video_id: str) -> dict | None:
    try:
        p = _cache_path(video_id)
        return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None
    except Exception:
        return None


def save_cached(notes: dict) -> None:
    try:
        _cache_path(notes["video_id"]).write_text(json.dumps(notes, ensure_ascii=False), encoding="utf-8")
    except OSError:
        pass


def make_notes(link: str, keys: list[str] | None = None, status=None) -> dict:
    """Notes for one video link. status(stage: str, fraction: float)."""
    status = status or (lambda *_: None)
    keys = get_keys() if keys is None else keys
    vid = parse_video_id(link)
    if not vid:
        raise NotesError("That doesn't look like a YouTube video link.")
    hit = load_cached(vid)
    if hit:
        status("Done", 1.0)
        return hit
    if not keys and not get_groq_keys():
        raise NotesError("The notes reader is not switched on yet. (App owner: add the free GEMINI_API_KEY "
                         "in the app's secrets — see README.)")
    status("Opening the lecture", 0.02)
    info = get_video_info(vid)
    if info.get("is_live"):
        raise NotesError("This is a live stream. Make notes after the live class has ended.")
    if (info.get("duration") or 0) > MAX_SECONDS:
        raise NotesError(f"This lecture is {fmt_time(info['duration'])} long. Lectures of up to 2 hours can be used.")

    shots = None
    tmp = Path(tempfile.mkdtemp(prefix="l2n_"))
    try:
        try:
            path = download_video(vid, tmp, info.get("duration") or 0,
                                  progress=lambda f: status("Getting the video (best quality, 240p–720p)", 0.03 + 0.3 * f))
            status("Stopping on every written page", 0.35)
            shots = board_snapshots(path, progress=lambda f: status("Stopping on every written page", 0.35 + 0.2 * f))
        except NotesError:
            raise
        except Exception:
            shots = None          # YouTube refused the download here: let the model watch it instead
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    try:
        if shots:
            notes = notes_from_snapshots(info, shots, keys, status)
        else:
            notes = notes_from_youtube(info, keys, status)
    except QuotaError:
        raise NotesError("Today's free reading limit is used up. Please try again tomorrow.")
    if not notes["items"]:
        raise NotesError("No questions or solutions were found written in this video.")
    notes.update({"video_id": vid, "video_title": info.get("title") or "", "duration": info.get("duration") or 0,
                  "max_height": info.get("max_height") or 0, "pages": len(shots or []),
                  "created": datetime.now().strftime("%d %b %Y, %H:%M")})
    save_cached(notes)
    status("Done", 1.0)
    return notes


# --------------------------------------------------------------------------- background jobs

@dataclass
class Job:
    links: list[str]
    id: str = field(default_factory=lambda: uuid.uuid4().hex[:10])
    fractions: list = field(default_factory=list)
    stages: list = field(default_factory=list)
    results: list = field(default_factory=list)     # notes or None, in link order
    errors: list = field(default_factory=list)      # (link number, link, message)
    done: bool = False
    started: float = field(default_factory=time.time)
    finished_at: float = 0.0
    _lock: threading.Lock = field(default_factory=threading.Lock, repr=False, compare=False)

    def __post_init__(self):
        self.fractions = [0.0] * len(self.links)
        self.stages = ["Waiting"] * len(self.links)
        self.results = [None] * len(self.links)
        self._left = len(self.links)
        self._closed: set[int] = set()

    def _set(self, i: int, stage: str, frac: float) -> None:
        if self.fractions[i] < 1.0:
            self.stages[i] = stage
            self.fractions[i] = max(self.fractions[i], min(max(frac, 0.0), 1.0))

    def _finish(self, i: int, result: dict | None, error: str | None) -> None:
        with self._lock:
            if i in self._closed:
                return
            self._closed.add(i)
            self.results[i] = result
            if error:
                self.errors.append((i + 1, self.links[i], error))
            self.stages[i], self.fractions[i] = "Done", 1.0
            self._left -= 1
            if self._left <= 0:
                self.errors.sort()
                self.finished_at = time.time()
                self.done = True

    @property
    def total(self) -> int:
        return len(self.links)

    @property
    def progress(self) -> float:
        return sum(self.fractions) / max(1, self.total)

    @property
    def finished(self) -> int:
        return sum(1 for f in self.fractions if f >= 1.0)

    @property
    def stage(self) -> str:
        active = [s for s, f in zip(self.stages, self.fractions) if 0 < f < 1]
        if active:
            return active[0]
        waiting = [s for s, f in zip(self.stages, self.fractions) if f < 1]
        return "Done" if self.done else (waiting[0] if waiting else "Starting")


JOBS: dict[str, Job] = {}   # all running/finished jobs, so a page refresh can find its job again

# One queue for every user of the app. A few lectures are worked on at a time (so the server never runs
# out of memory), users take fair turns (one lecture each, round and round), and a lecture asked for by
# many users at once is made only once — everyone waiting for it gets the same notes.


class _Flight:
    """One video on its way to notes, shared by every (job, link number) that asked for it."""

    def __init__(self, key: str, link: str, keys):
        self.key, self.link, self.keys = key, link, keys
        self.subs: list[tuple[Job, int]] = []

    def status(self, stage: str, frac: float) -> None:
        for job, i in list(self.subs):
            job._set(i, stage, frac)


_q_lock = threading.Condition()
_flights: dict[str, _Flight] = {}          # video → queued or running flight
_pending: dict[str, list[_Flight]] = {}    # job id → its flights that have not started yet
_turns: list[str] = []                     # job ids, round-robin order
_workers: list[threading.Thread] = []


def _worker_count() -> int:
    try:
        return max(1, int(os.environ.get("NOTES_WORKERS") or WORKERS))
    except ValueError:
        return WORKERS


def _renumber() -> None:
    """Tell every waiting link its place in the line (caller holds _q_lock)."""
    queues = {jid: list(_pending.get(jid) or []) for jid in _turns}
    order, place = list(_turns), 0
    while order:
        jid = order.pop(0)
        if queues[jid]:
            place += 1
            f = queues[jid].pop(0)
            for job, i in f.subs:
                if job.fractions[i] == 0:
                    job.stages[i] = f"Waiting in line · number {place}"
            if queues[jid]:
                order.append(jid)


def _next_flight() -> _Flight | None:
    while _turns:
        jid = _turns.pop(0)
        q = _pending.get(jid)
        if not q:
            _pending.pop(jid, None)
            continue
        f = q.pop(0)
        if q:
            _turns.append(jid)       # back of the line for its next lecture
        else:
            _pending.pop(jid, None)
        _renumber()
        return f
    return None


def _run_flight(f: _Flight) -> None:
    result = error = None
    try:
        result = make_notes(f.link, f.keys, f.status)
    except NotesError as e:
        error = str(e)
    except Exception as e:  # never leave the student with a frozen screen
        error = f"Something went wrong with this video ({type(e).__name__}). Please try again."
    with _q_lock:
        _flights.pop(f.key, None)
        subs = list(f.subs)
    for job, i in subs:
        job._finish(i, result, error)


def _worker() -> None:
    while True:
        with _q_lock:
            f = _next_flight()
            while f is None:
                _q_lock.wait()
                f = _next_flight()
        _run_flight(f)


def _prune_jobs() -> None:
    old = time.time() - JOB_KEEP_SECONDS
    for jid in [j for j, job in JOBS.items() if job.done and job.finished_at and job.finished_at < old]:
        JOBS.pop(jid, None)


def start_job(job: Job, keys: list[str] | None = None) -> None:
    """Put a job's links in the shared line. Cached lectures are filled in at once."""
    _prune_jobs()
    cached = []
    with _q_lock:
        for i, link in enumerate(job.links):
            vid = parse_video_id(link)
            hit = load_cached(vid) if vid else None
            if hit:
                cached.append((i, hit))
                continue
            key = vid or f"link:{link}"
            f = _flights.get(key)
            if f is None:
                f = _flights[key] = _Flight(key, link, keys)
                _pending.setdefault(job.id, []).append(f)
            f.subs.append((job, i))
            job.stages[i] = "Waiting in line"
        if _pending.get(job.id) and job.id not in _turns:
            _turns.append(job.id)
        _renumber()
        _workers[:] = [t for t in _workers if t.is_alive()]
        while len(_workers) < _worker_count():
            t = threading.Thread(target=_worker, daemon=True, name=f"notes-worker-{len(_workers) + 1}")
            t.start()
            _workers.append(t)
        _q_lock.notify_all()
    for i, hit in cached:
        job._finish(i, hit, None)
    if not job.links:
        job.done = True


def run_job(job: Job, keys: list[str] | None = None) -> None:
    """start_job, then wait until every link is finished."""
    start_job(job, keys)
    while not job.done:
        _sleep(0.02)


_sleep = time.sleep


def combine(results: list[dict]) -> dict | None:
    """All lectures in one document, in the order the links were given."""
    done = [r for r in results if r]
    if not done:
        return None
    if len(done) == 1:
        return dict(done[0], lectures=1)
    items: list[dict] = []
    for n, r in enumerate(done, 1):
        items.append({"type": "lecture", "text": f"Lecture {n} · {r['title']}"})
        items += r["items"]
    return {"title": "Lecture notes", "items": items, "lectures": len(done),
            "duration": sum(r.get("duration") or 0 for r in done),
            "created": datetime.now().strftime("%d %b %Y, %H:%M")}


# --------------------------------------------------------------------------- export

def grouped(items: list[dict]):
    """Yield (kind, payload) blocks: questions get their solution and answer attached."""
    qn, i = 0, 0
    while i < len(items):
        it = items[i]
        if it["type"] == "question":
            qn += 1
            steps, answers, j = [], [], i + 1
            while j < len(items) and items[j]["type"] in ("step", "answer", "formula") and not (
                    items[j]["type"] == "formula" and answers):
                (answers if items[j]["type"] == "answer" else steps).append(items[j])
                j += 1
            yield "question", {"n": qn, "text": it["text"], "steps": steps, "answers": answers}
            i = j
            continue
        if it["type"] == "step":
            steps, j = [], i
            while j < len(items) and items[j]["type"] == "step":
                steps.append(items[j])
                j += 1
            yield "steps", {"steps": steps}
            i = j
            continue
        yield it["type"], it
        i += 1


def safe_filename(title: str) -> str:
    name = re.sub(r"[^\w\s-]", "", title or "notes").strip()
    return (re.sub(r"\s+", " ", name)[:70] or "notes").strip()


def to_text(notes: dict) -> str:
    out = [notes["title"], "=" * min(len(notes["title"]), 70), ""]
    for kind, b in grouped(notes["items"]):
        if kind == "lecture":
            out += ["", "#" * 60, b["text"], "#" * 60]
        elif kind == "heading":
            out += ["", b["text"].upper(), "-" * min(len(b["text"]), 70)]
        elif kind == "question":
            out += ["", f"Q{b['n']}. {b['text']}"]
            if b["steps"]:
                out.append("Solution:")
                out += [f"   {s['text']}" for s in b["steps"]]
            out += [f"Answer: {a['text']}" for a in b["answers"]]
        elif kind == "steps":
            out += [f"   {s['text']}" for s in b["steps"]]
        elif kind == "answer":
            out.append(f"Answer: {b['text']}")
        elif kind == "definition":
            out.append(f"Definition: {b['text']}")
        elif kind == "formula":
            out.append(f"    {b['text']}")
        else:
            out.append(b["text"])
    return "\n".join(out).strip() + "\n"


def _font(name: str) -> str:
    for p in (FONTS / name, Path("/usr/share/fonts/truetype/dejavu") / name):
        if p.exists():
            return str(p)
    raise NotesError("PDF font files are missing (fonts folder).")


def to_pdf(notes: dict) -> bytes:
    from fpdf import FPDF

    ink, muted, accent, soft = (24, 24, 27), (100, 100, 108), (190, 92, 58), (246, 243, 238)

    class PDF(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("dv", "", 8)
            self.set_text_color(*muted)
            self.cell(0, 6, f"{notes['title'][:80]}  ·  page {self.page_no()}", align="C")
            self.set_y(-7)
            self.set_font("dv", "", 7.5)
            self.cell(0, 4, CREDIT, align="R")

    pdf = PDF(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, 18)
    pdf.add_font("dv", "", _font("DejaVuSans.ttf"))
    pdf.add_font("dv", "B", _font("DejaVuSans-Bold.ttf"))
    pdf.add_font("dv", "I", _font("DejaVuSans-Oblique.ttf"))
    pdf.add_page()
    w = pdf.epw

    def para(text, size=10.5, style="", color=ink, indent=0, h=5.6, fill=None, align="L"):
        pdf.set_font("dv", style, size)
        pdf.set_text_color(*color)
        pdf.set_x(pdf.l_margin + indent)
        if fill:
            pdf.set_fill_color(*fill)
        pdf.multi_cell(w - indent, h, text, fill=bool(fill), align=align, new_x="LMARGIN", new_y="NEXT")

    para(notes["title"], 19, "B", h=9)
    sub = " · ".join(x for x in [notes.get("created", ""),
                                 f"{notes['lectures']} lectures" if notes.get("lectures", 1) > 1 else ""] if x)
    if sub:
        para(sub, 9, color=muted)
    pdf.ln(2)
    pdf.set_draw_color(*accent)
    pdf.set_line_width(0.6)
    pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + 30, pdf.get_y())
    pdf.ln(5)

    for kind, b in grouped(notes["items"]):
        if pdf.get_y() > pdf.h - 40 and kind in ("heading", "question", "lecture"):
            pdf.add_page()
        if kind == "lecture":
            pdf.ln(4)
            para(b["text"], 15, "B", h=8)
            pdf.set_draw_color(220, 216, 210)
            pdf.set_line_width(0.3)
            pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.l_margin + w, pdf.get_y() + 1)
            pdf.ln(4)
        elif kind == "heading":
            pdf.ln(3)
            para(b["text"], 13.5, "B", color=accent, h=7)
            pdf.ln(1.5)
        elif kind == "question":
            pdf.ln(2.5)
            para(f"Q{b['n']}.  {b['text']}", 11, "B", fill=soft, h=6.2)
            if b["steps"]:
                pdf.ln(1)
                para("Solution", 9, "I", color=muted)
                for s in b["steps"]:
                    para(s["text"], 10.5, indent=6, h=6)
            for a in b["answers"]:
                pdf.ln(1)
                para(f"Answer:  {a['text']}", 10.5, "B", color=accent, indent=2)
            pdf.ln(2)
        elif kind == "steps":
            for s in b["steps"]:
                para(s["text"], 10.5, indent=6, h=6)
        elif kind == "answer":
            para(f"Answer:  {b['text']}", 10.5, "B", color=accent, indent=2)
        elif kind == "definition":
            para(f"Definition:  {b['text']}", 10.5)
            pdf.ln(1)
        elif kind == "formula":
            pdf.ln(1)
            para(b["text"], 11, fill=soft, align="C", h=7)
            pdf.ln(1)
        else:
            para(b["text"])
            pdf.ln(1.5)
    return bytes(pdf.output())


def to_docx(notes: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    accent, muted = RGBColor(0xBE, 0x5C, 0x3A), RGBColor(0x64, 0x64, 0x6C)
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
        f = s.footer.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fr = f.add_run(CREDIT)
        fr.font.size, fr.font.color.rgb = Pt(8), muted
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(11)
    base.paragraph_format.space_after = Pt(4)

    def shade(par, color="F6F3EE"):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        par._p.get_or_add_pPr().append(shd)

    def add(text, bold=False, italic=False, color=None, size=None, indent=None, math=False, align=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if color:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
        if math:
            r.font.name = "Cambria Math"
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        if align:
            p.alignment = align
        return p

    t = doc.add_heading(notes["title"], level=0)
    for r in t.runs:
        r.font.color.rgb = RGBColor(0x18, 0x18, 0x1B)
    for kind, b in grouped(notes["items"]):
        if kind == "lecture":
            doc.add_heading(b["text"], level=1)
        elif kind == "heading":
            h = doc.add_heading(b["text"], level=2)
            for r in h.runs:
                r.font.color.rgb = accent
        elif kind == "question":
            p = add(f"Q{b['n']}.  {b['text']}", bold=True)
            shade(p)
            p.paragraph_format.space_before = Pt(8)
            if b["steps"]:
                add("Solution", italic=True, color=muted, size=9.5)
                for s in b["steps"]:
                    add(s["text"], indent=0.8, math=True).paragraph_format.space_after = Pt(1)
            for a in b["answers"]:
                add(f"Answer:  {a['text']}", bold=True, color=accent, indent=0.3)
        elif kind == "steps":
            for s in b["steps"]:
                add(s["text"], indent=0.8, math=True).paragraph_format.space_after = Pt(1)
        elif kind == "answer":
            add(f"Answer:  {b['text']}", bold=True, color=accent, indent=0.3)
        elif kind == "definition":
            p = doc.add_paragraph()
            p.add_run("Definition:  ").bold = True
            p.add_run(b["text"])
        elif kind == "formula":
            shade(add(b["text"], math=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER))
        else:
            add(b["text"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
